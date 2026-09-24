import asyncio
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Union

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.logger.logger_wrapper import LoggerWrapper
from app.utils import Utils

logger = LoggerWrapper()

class DocWriterService:

    def __init__(self) -> None:
        config = Utils.get_config_file()
        reports_path = config.get("doc_writer", {}).get("output_dir", None)

        self.company_name: Optional[str] = None
        self.output_dir = Path(__file__).parent.parent.parent / reports_path
        self.output_dir.mkdir(parents=True, exist_ok=True)

        logger(f"[DocWriterService_c:output_dir_v]: {self.output_dir}")

    def set_company_name(self, company_name: Optional[str]):
        self.company_name = company_name

    async def save_answers(self,
        pairs: List[tuple[Optional[str], str]],
        metadata: Optional[Dict[str, str]] = None) -> Optional[Path]:

        if not self.company_name:
            self.company_name = "DEFAULT_COMPANY"

        try:
            file_path = self._resolve_path()
            document = self._build_document(pairs, metadata)

            await asyncio.to_thread(document.save, file_path)

            logger(f"[DocWriterService_c:save_answers_f:output_dir_v]: {file_path}")
            return file_path
        except Exception as e:
            logger(f"[DocWriterService_c:save_answers_f:file_path_err]: {e}")
            return None

    def _build_document(self,
        pairs: List[tuple[Optional[str], str]],
        metadata: Optional[Dict[str, str]]) -> Document:

        document = Document()

        self._add_title(document, self.company_name)
        if metadata:
            self._add_metadata(document, metadata)

        for index, (question, answer) in enumerate(pairs, start=1):
            if question:
                self._add_question(document, index, question)
            self._add_answer(document, answer)

        return document

    def _resolve_path(self) -> Path:

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.company_name = f"{self.company_name}_{timestamp}"

        if not self.company_name.lower().endswith(".docx"):
            self.company_name += ".docx"

        return self.output_dir / self.company_name

    @staticmethod
    def _add_title(document: Document, filename: Optional[str]) -> None:
        title = document.add_heading(filename, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _add_metadata(document: Document, metadata: Dict[str, str]) -> None:
        for key, value in metadata.items():
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{key}: {value}")
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    @staticmethod
    def _add_question(document: Document, index: int, question: str) -> None:
        heading = document.add_heading(f"Q{index}. {question}", level=2)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    @staticmethod
    def _add_answer(document: Document, answer: str) -> None:
        for line in answer.split("\n"):
            if line.strip():
                document.add_paragraph(line)