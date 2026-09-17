import pytest
from pathlib import Path
from unittest.mock import patch
import sys

project_root = Path(__file__).parent.parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

from app.service.doc_service import DocWriterService
from docx import Document


@pytest.fixture
def mock_config(tmp_path):
    return {"doc_writer": {"output_dir": str(tmp_path)}}


@pytest.fixture
def writer(mock_config, tmp_path):
    with patch("app.service.doc_service.Utils.get_config_file", return_value=mock_config):
        yield DocWriterService()


@pytest.mark.asyncio
class TestDocWriterService:

    # async def test_save_answer_creates_file(self, writer, tmp_path):
    #     path = await writer.save_answer(
    #         question="What is Python?",
    #         answer="A programming language.",
    #     )
    #
    #     assert path is not None
    #     assert path.exists()
    #     assert path.suffix == ".docx"
    #     assert path.parent == tmp_path
    #
    # async def test_save_answer_content(self, writer):
    #     path = await writer.save_answer(
    #         question="What is Python?",
    #         answer="A programming language.",
    #         filename="test_content",
    #     )
    #
    #     document = Document(path)
    #     text = "\n".join(p.text for p in document.paragraphs)
    #
    #     assert "LLM Answers" in text
    #     assert "What is Python?" in text
    #     assert "A programming language." in text

    async def test_save_multiple_answers(self, writer):
        pairs = [
            ("Question 1", "Answer 1"),
            ("Question 2", "Answer 2"),
            (None, "Unlabeled answer"),
        ]
        path = await writer.save_answers(pairs=pairs, filename="multi")

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)

        assert "Question 1" in text
        assert "Answer 1" in text
        assert "Question 2" in text
        assert "Answer 2" in text
        assert "Unlabeled answer" in text

    # async def test_metadata_added(self, writer):
    #     path = await writer.save_answer(
    #         answer="Answer",
    #         filename="meta",
    #         metadata={"Model": "gpt-4", "Date": "2025-01-15"},
    #     )
    #
    #     document = Document(path)
    #     text = "\n".join(p.text for p in document.paragraphs)
    #
    #     assert "Model: gpt-4" in text
    #     assert "Date: 2025-01-15" in text
    #
    # async def test_default_filename_with_timestamp(self, writer):
    #     path = await writer.save_answer(answer="Answer")
    #
    #     assert path is not None
    #     assert path.name.startswith("llm_answers_")
    #     assert path.suffix == ".docx"
    #
    # async def test_docx_extension_added(self, writer):
    #     path = await writer.save_answer(answer="Answer", filename="no_ext")
    #     assert path.name == "no_ext.docx"
    #
    #     path2 = await writer.save_answer(answer="Answer", filename="with_ext.docx")
    #     assert path2.name == "with_ext.docx"
    #
    # async def test_multiline_answer(self, writer):
    #     answer = "Line 1\nLine 2\n\nLine 3"
    #     path = await writer.save_answer(answer=answer, filename="multiline")
    #
    #     document = Document(path)
    #     paragraphs = [p.text for p in document.paragraphs if p.text]
    #
    #     assert "Line 1" in paragraphs
    #     assert "Line 2" in paragraphs
    #     assert "Line 3" in paragraphs
    #
    # async def test_save_handles_error(self, writer, monkeypatch):
    #     """Test that save returns None on error."""
    #     def broken_save(self, *args, **kwargs):
    #         raise Exception("Disk full")
    #
    #     monkeypatch.setattr(Document, "save", broken_save)
    #
    #     path = await writer.save_answer(answer="Answer", filename="error")
    #     assert path is None